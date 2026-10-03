import io
from pathlib import PurePath

from docx import Document
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # CLAUDE.md rule 9


class UnreadableResume(ValueError):
    """The upload cannot be turned into text. The message is safe to show the user."""


def _pdf(data: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)


def _docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    # resumes built from Word templates often keep skills and dates in tables
    lines += ["\t".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
    return "\n".join(lines)


_EXTRACTORS = {".pdf": _pdf, ".docx": _docx, ".txt": lambda data: data.decode("utf-8", "replace")}


def extract_text(filename: str, data: bytes) -> str:
    """Resume bytes to plain text. Works from memory; the file is never written to disk."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise UnreadableResume("File is larger than 5 MB.")
    extractor = _EXTRACTORS.get(PurePath(filename).suffix.lower())
    if extractor is None:
        raise UnreadableResume("Upload a PDF, DOCX or TXT file.")
    try:
        text = extractor(data).strip()
    except Exception:  # noqa: BLE001 - pypdf and python-docx raise many unrelated types
        raise UnreadableResume("The file could not be read. Is it a valid document?") from None
    if len(text) < 200:
        # a scanned PDF has pages but no text layer
        raise UnreadableResume("No text found. If this is a scanned PDF, export a text-based one.")
    return text
