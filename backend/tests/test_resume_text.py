import io
from pathlib import Path

import pytest
from docx import Document

from app import sample
from app.core.resume_text import MAX_UPLOAD_BYTES, UnreadableResume, extract_text

RESUME = sample.text("resume.txt")


def test_pdf():
    data = (Path(__file__).parent / "fixtures" / "sample_resume.pdf").read_bytes()
    text = extract_text("resume.pdf", data)
    assert "Kestrel Data Labs" in text and "1,200 internal support documents" in text


def test_docx_including_tables():
    doc = Document()
    for line in RESUME.splitlines():
        doc.add_paragraph(line)
    doc.add_table(rows=1, cols=2).rows[0].cells[1].text = "in-a-table"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("Resume.DOCX", buf.getvalue())
    assert "Kestrel Data Labs" in text and "in-a-table" in text


@pytest.mark.parametrize(
    ("name", "data"),
    [
        ("resume.png", b"x" * 500),
        ("resume.txt", b"x" * (MAX_UPLOAD_BYTES + 1)),
        ("resume.pdf", b"not a pdf at all" * 20),
        ("resume.txt", b"too short"),
    ],
    ids=["wrong type", "too large", "corrupt", "no text"],
)
def test_rejects_unusable_uploads(name, data):
    with pytest.raises(UnreadableResume):
        extract_text(name, data)
